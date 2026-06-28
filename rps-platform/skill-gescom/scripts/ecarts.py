# -*- coding: utf-8 -*-
# Excel des ecarts vs rapport officiel. Necessite rep_officiel.json (dict num->[GASOIL,SUPER]).
# Lit prises.json + report_data.json + missing/horsservice (cwd). Sortie ECARTS-VENTES-<date>.xlsx dans RPS_DAYDIR.
import json,os
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter

if not os.path.exists('rep_officiel.json'):
    print("ECARTS: rep_officiel.json absent -> Excel des ecarts non genere (fournir le rapport officiel)."); raise SystemExit(0)
DAYDIR=os.environ["RPS_DAYDIR"]
d=json.load(open('report_data.json')); DAY=d.get('day','')
DAYLBL="%s/%s/20%s"%(DAY[0:2],DAY[2:4],DAY[4:6]) if len(DAY)==6 else DAY
REP=json.load(open('rep_officiel.json'))
recs=json.load(open('prises.json'))
nom={s[0]:s[1] for s in d['stations']}
missing={m['num']:m['nom'] for m in (json.load(open('missing.json')) if os.path.exists('missing.json') else [])}
hs={m['num'] for m in (json.load(open('horsservice.json')) if os.path.exists('horsservice.json') else [])}
mine=defaultdict(lambda:{'GASOIL':0.0,'SUPER':0.0})
for r in recs: mine[r['snum']][r['produit']]+=r['qte']
def cause(num,dg,ds):
    if num in missing: return "Station non transmise - a relancer"
    if num in hs: return "Hors service"
    if abs(dg)<=1 and abs(ds)<=1: return "Concordant"
    return "A verifier"
wb=Workbook(); ws=wb.active; ws.title="Ecarts %s"%DAYLBL.replace('/','-')
NAVY="1A3C6E"; LRED="FBEAEA"; LGREEN="E6F4EA"
thin=Side(style='thin',color="CCCCCC"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
def setc(r,c,v,**kw):
    cell=ws.cell(row=r,column=c,value=v); cell.border=border
    cell.font=Font(name='Arial',size=kw.get('size',10),bold=kw.get('bold',False),color=kw.get('color','000000'))
    if kw.get('fill'): cell.fill=PatternFill('solid',fgColor=kw['fill'])
    cell.alignment=Alignment(horizontal=kw.get('h','left'),vertical='center',wrap_text=kw.get('wrap',False))
    if 'fmt' in kw: cell.number_format=kw['fmt']
    return cell
ws.merge_cells('A1:H1'); setc(1,1,"RPS — Ecarts ventes GESCOM vs Rapport officiel — %s"%DAYLBL,bold=True,size=13,color="FFFFFF",fill=NAVY,h='center')
ws.merge_cells('A2:H2'); setc(2,1,"Quantites en litres. Aucune valeur GESCOM modifiee — ecarts signales pour relance.",size=8,color="555555")
ws.row_dimensions[1].height=22
hdr=["N°","Station","Officiel GASOIL","Officiel SUPER","GESCOM GASOIL","GESCOM SUPER","Ecart GASOIL","Ecart SUPER","Cause"]
hr=4
for j,h in enumerate(hdr,1): setc(hr,j,h,bold=True,color="FFFFFF",fill=NAVY,h='center',wrap=True)
ws.row_dimensions[hr].height=30
rowsd=[]
for num,(rg,rs) in REP.items():
    mg=round(mine.get(num,{}).get('GASOIL',0.0),2); ms=round(mine.get(num,{}).get('SUPER',0.0),2)
    rowsd.append((num,rg,rs,mg,ms))
rowsd.sort(key=lambda x:(-(abs(x[3]-x[1])+abs(x[4]-x[2])), x[0]))
r=hr+1; first=r
for num,rg,rs,mg,ms in rowsd:
    st=nom.get(num) or missing.get(num,num)
    dg=mg-rg; ds=ms-rs; ecr=(abs(dg)>1 or abs(ds)>1); fill=LRED if ecr else LGREEN
    setc(r,1,num,h='center'); setc(r,2,str(st))
    setc(r,3,rg,h='right',fmt='#,##0.00;(#,##0.00);-'); setc(r,4,rs,h='right',fmt='#,##0.00;(#,##0.00);-')
    setc(r,5,mg,h='right',fmt='#,##0.00;(#,##0.00);-'); setc(r,6,ms,h='right',fmt='#,##0.00;(#,##0.00);-')
    setc(r,7,"=E%d-C%d"%(r,r),h='right',fmt='#,##0.00;(#,##0.00);-',fill=fill,bold=ecr)
    setc(r,8,"=F%d-D%d"%(r,r),h='right',fmt='#,##0.00;(#,##0.00);-',fill=fill,bold=ecr)
    setc(r,9,cause(num,dg,ds),size=8.5,fill=(LRED if ecr else None)); r+=1
last=r-1
setc(r,2,"TOTAL",bold=True,fill=NAVY,color="FFFFFF")
for col in (3,4,5,6,7,8):
    L=get_column_letter(col); setc(r,col,"=SUM(%s%d:%s%d)"%(L,first,L,last),h='right',bold=True,fmt='#,##0.00;(#,##0.00);-',fill=NAVY,color="FFFFFF")
setc(r,9,"Ecarts = stations non transmises + a verifier",size=8.5,bold=True,fill=NAVY,color="FFFFFF",wrap=True)
for j,w in enumerate([6,30,16,16,16,16,15,15,40],1): ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes="A5"; ws.auto_filter.ref="A%d:I%d"%(hr,last)
OUT=os.path.join(DAYDIR,"ECARTS-VENTES-%s.xlsx"%DAYLBL.replace('/','-'))
wb.save(OUT); print("ECARTS OK ->",OUT)
